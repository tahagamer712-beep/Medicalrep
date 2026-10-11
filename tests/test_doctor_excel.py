"""Doctor Excel: export, template, preview, apply, rollback, protection."""
import pytest
from openpyxl import Workbook, load_workbook

from app.core.errors import AlreadyExists, PermissionDenied, ValidationError
from app.databases import manager as dbm
from app.databases import office_models as om
from app.services import doctor_excel as dx
from app.services import doctors as docs
from app.services import products as prod


@pytest.fixture()
def catalog(seed):
    d = seed["data_dir"]
    u = prod.create_unit(d, seed["supervisor1"], office_id=seed["office1"],
                         name="علبة")
    prod.create_product(d, seed["supervisor1"], office_id=seed["office1"],
                        trade_name="Panadol", unit_id=u["id"])
    prod.create_product(d, seed["supervisor1"], office_id=seed["office1"],
                        trade_name="Brufen", unit_id=u["id"])
    out = docs.create_doctor(
        d, seed["supervisor1"], office_id=seed["office1"],
        name="د. أحمد", area="المنصور", specialty="باطنية",
        doctor_class="A", assigned_rep_id=seed["rep1_o1"])
    with dbm.office_session(d, seed["office1"]) as s:
        code = s.get(om.Doctor, out["id"]).external_code
    return code


def _write_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Doctors"
    ws.append(dx.COLUMNS)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def _row(code="", ver="", name="د. جديد", area="المنصور", sub="",
         pharm="", spec="باطنية", cls="A", tgt="Panadol; Brufen",
         info="", rdays="0,2", ndays="", rtime="09:00", notes="",
         dealer="", status="Active"):
    return [code, ver, name, area, sub, pharm, spec, cls, tgt, info,
            rdays, ndays, rtime, notes, dealer, status]


def test_export_contains_doctors(seed, catalog, tmp_path):
    d = seed["data_dir"]
    p = dx.export_doctors(d, seed["supervisor1"], office_id=seed["office1"],
                          out_path=tmp_path / "docs.xlsx")
    wb = load_workbook(p)
    ws = wb["Doctors"]
    assert ws.max_row == 2  # header + 1 doctor
    assert ws.cell(1, 1).value == "Doctor Code"


def test_rep_exports_only_assigned(seed, catalog, tmp_path):
    d = seed["data_dir"]
    p = dx.export_doctors(d, seed["rep1_o1"], office_id=seed["office1"],
                          out_path=tmp_path / "r1.xlsx")
    assert load_workbook(p)["Doctors"].max_row == 2
    p2 = dx.export_doctors(d, seed["rep2_o1"], office_id=seed["office1"],
                           out_path=tmp_path / "r2.xlsx")
    assert load_workbook(p2)["Doctors"].max_row == 1  # header only


def test_preview_and_apply_new(seed, catalog, tmp_path):
    d = seed["data_dir"]
    f = _write_xlsx(tmp_path / "imp.xlsx",
                    [_row(name="د. جديد", tgt="Panadol")])
    pv = dx.preview_import(d, seed["supervisor1"], office_id=seed["office1"],
                           file_path=f)
    assert pv["can_apply"] and pv["new"] == 1 and pv["updated"] == 0
    # needs explicit confirmation
    with pytest.raises(ValidationError):
        dx.apply_import(d, seed["supervisor1"], office_id=seed["office1"],
                        file_path=f, confirm=False)
    out = dx.apply_import(d, seed["supervisor1"], office_id=seed["office1"],
                          file_path=f, confirm=True)
    assert out["added"] == 1
    found = docs.search_doctors(d, seed["supervisor1"],
                                office_id=seed["office1"], name="د. جديد")
    assert len(found) == 1
    # same file twice -> rejected
    with pytest.raises(AlreadyExists):
        dx.apply_import(d, seed["supervisor1"], office_id=seed["office1"],
                        file_path=f, confirm=True)


def test_update_roundtrip(seed, catalog, tmp_path):
    d = seed["data_dir"]
    exp = tmp_path / "exp.xlsx"
    dx.export_doctors(d, seed["supervisor1"], office_id=seed["office1"],
                      out_path=exp)
    wb = load_workbook(exp)
    ws = wb["Doctors"]
    code = ws.cell(2, 1).value
    ver = ws.cell(2, 2).value
    ws.cell(2, 4).value = "اليرموك"  # Area column
    wb.save(exp)
    pv = dx.preview_import(d, seed["supervisor1"], office_id=seed["office1"],
                           file_path=exp)
    assert pv["updated"] == 1 and pv["can_apply"]
    dx.apply_import(d, seed["supervisor1"], office_id=seed["office1"],
                    file_path=exp, confirm=True)
    with dbm.office_session(d, seed["office1"]) as s:
        doc = s.query(om.Doctor).filter_by(external_code=code).one()
        assert doc.area == "اليرموك" and doc.version == ver + 1


def test_conflict_detected(seed, catalog, tmp_path):
    d = seed["data_dir"]
    exp = tmp_path / "exp.xlsx"
    dx.export_doctors(d, seed["supervisor1"], office_id=seed["office1"],
                      out_path=exp)
    # someone edits the doctor in the system after the export
    with dbm.office_session(d, seed["office1"]) as s:
        doc = s.query(om.Doctor).first()
        did = doc.id
    v = docs.get_doctor(d, seed["supervisor1"], office_id=seed["office1"],
                        doctor_id=did)["version"]
    docs.update_doctor(d, seed["supervisor1"], office_id=seed["office1"],
                       doctor_id=did, expected_version=v, area="الكاظمية")
    pv = dx.preview_import(d, seed["supervisor1"], office_id=seed["office1"],
                           file_path=exp)
    assert len(pv["conflicts"]) == 1 and not pv["can_apply"]
    with pytest.raises(ValidationError):
        dx.apply_import(d, seed["supervisor1"], office_id=seed["office1"],
                        file_path=exp, confirm=True)
    # system edit was NOT overwritten
    got = docs.get_doctor(d, seed["supervisor1"], office_id=seed["office1"],
                          doctor_id=did)
    assert got["area"] == "الكاظمية"


def test_row_errors_listed_and_nothing_applied(seed, catalog, tmp_path):
    d = seed["data_dir"]
    f = _write_xlsx(tmp_path / "bad.xlsx", [
        _row(name="", cls="X", status="Weird", tgt="NoSuchDrug"),
        _row(name="د. سليم", tgt="Panadol"),
    ])
    pv = dx.preview_import(d, seed["supervisor1"], office_id=seed["office1"],
                           file_path=f)
    assert len(pv["errors"]) == 1
    assert pv["errors"][0]["row"] == 2
    assert not pv["can_apply"]
    with pytest.raises(ValidationError):
        dx.apply_import(d, seed["supervisor1"], office_id=seed["office1"],
                        file_path=f, confirm=True)
    # nothing was applied, not even the valid row
    assert docs.search_doctors(d, seed["supervisor1"],
                               office_id=seed["office1"],
                               name="د. سليم") == []


def test_missing_column_rejected(seed, tmp_path):
    d = seed["data_dir"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Doctors"
    ws.append(["Dr Name"])  # wrong header
    ws.append(["د. س"])
    f = tmp_path / "bad.xlsx"
    wb.save(f)
    with pytest.raises(ValidationError):
        dx.preview_import(d, seed["supervisor1"], office_id=seed["office1"],
                          file_path=f)


def test_empty_file_rejected(seed, tmp_path):
    d = seed["data_dir"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Doctors"
    ws.append(dx.COLUMNS)
    f = tmp_path / "empty.xlsx"
    wb.save(f)
    with pytest.raises(ValidationError):
        dx.preview_import(d, seed["supervisor1"], office_id=seed["office1"],
                          file_path=f)


def test_rollback_import(seed, catalog, tmp_path):
    d = seed["data_dir"]
    f = _write_xlsx(tmp_path / "imp.xlsx",
                    [_row(name="د. جديد", area="المنصور")])
    out = dx.apply_import(d, seed["supervisor1"], office_id=seed["office1"],
                          file_path=f, confirm=True)
    rb = dx.rollback_import(d, seed["supervisor1"], office_id=seed["office1"],
                            import_id=out["import_id"])
    assert rb["doctors_touched"] == 1
    got = docs.search_doctors(d, seed["supervisor1"],
                              office_id=seed["office1"], name="د. جديد")
    assert got[0]["status"] == "inactive"  # deactivated, not deleted
    with dbm.office_session(d, seed["office1"]) as s:
        log = s.get(om.ImportLog, out["import_id"])
        assert log.status == "rolled_back"


def test_rep_cannot_touch_others_doctors(seed, catalog, tmp_path):
    d = seed["data_dir"]
    exp = tmp_path / "exp.xlsx"
    dx.export_doctors(d, seed["supervisor1"], office_id=seed["office1"],
                      out_path=exp)
    pv = dx.preview_import(d, seed["rep2_o1"], office_id=seed["office1"],
                           file_path=exp)
    assert len(pv["errors"]) == 1  # out of rep2's scope
    assert "صلاحياتك" in pv["errors"][0]["reason"]
