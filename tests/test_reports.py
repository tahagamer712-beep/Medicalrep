"""Reports: daily/weekly/monthly + performance + Excel/CSV export."""
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.core.errors import PermissionDenied, ValidationError
from app.services import attendance as att
from app.services import leave as lv
from app.services import payroll as pr
from app.services import plans
from app.services import reports as rep
from app.services import visits as vsvc

DAY = "2026-10-12"  # Monday
MONDAY = "2026-10-12"


@pytest.fixture()
def activity(seed):
    d = seed["data_dir"]
    o1 = seed["office1"]
    doc = vsvc.add_doctor(d, seed["supervisor1"], office_id=o1, name="د. ر")
    # rep1: visit + attendance + plan
    vsvc.record_actual_visit(
        d, seed["rep1_o1"], office_id=o1, rep_user_id=seed["rep1_o1"],
        target_type="doctor", target_id=doc["id"], outcome="done",
        confirm=True,
        visited_at=datetime(2026, 10, 12, 7, 0, tzinfo=timezone.utc))
    att.check_in(d, seed["rep1_o1"], office_id=o1, rep_user_id=seed["rep1_o1"],
                 now=datetime(2026, 10, 12, 5, 0, tzinfo=timezone.utc))
    p = plans.create_plan(d, seed["rep1_o1"], office_id=o1, plan_type="weekly",
                          rep_user_id=seed["rep1_o1"], period=MONDAY,
                          payload={})
    plans.submit_plan(d, seed["rep1_o1"], office_id=o1, plan_type="weekly",
                      plan_id=p["plan_id"])
    # rep2: attendance only
    att.check_in(d, seed["rep2_o1"], office_id=o1, rep_user_id=seed["rep2_o1"],
                 now=datetime(2026, 10, 12, 4, 50, tzinfo=timezone.utc))
    return o1


def test_daily_report(seed, activity):
    d = seed["data_dir"]
    r = rep.daily_report(d, seed["supervisor1"], office_id=activity, day=DAY)
    assert r["totals"]["visits"] == 1
    assert r["totals"]["present"] == 2
    by_rep = {x["rep_id"]: x for x in r["rows"]}
    assert by_rep[seed["rep1_o1"]]["visits"] == 1
    assert by_rep[seed["rep1_o1"]]["weekly_plan"] == "submitted"
    assert by_rep[seed["rep2_o1"]]["visits"] == 0


def test_rep_sees_only_own_daily(seed, activity):
    d = seed["data_dir"]
    r = rep.daily_report(d, seed["rep1_o1"], office_id=activity, day=DAY)
    assert [x["rep_id"] for x in r["rows"]] == [seed["rep1_o1"]]


def test_weekly_report(seed, activity):
    d = seed["data_dir"]
    r = rep.weekly_report(d, seed["manager1"], office_id=activity,
                          week_start=MONDAY)
    by_rep = {x["rep_id"]: x for x in r["rows"]}
    assert by_rep[seed["rep1_o1"]]["visits"] == 1
    assert by_rep[seed["rep1_o1"]]["weekly_plan"] == "submitted"
    assert by_rep[seed["rep1_o1"]]["days_present"] == 1
    assert r["totals"]["visits"] == 1


def test_weekly_requires_monday(seed, activity):
    d = seed["data_dir"]
    with pytest.raises(ValidationError):
        rep.weekly_report(d, seed["supervisor1"], office_id=activity,
                          week_start="2026-10-13")  # Tuesday


def test_monthly_report_with_payroll(seed, activity):
    d = seed["data_dir"]
    p = pr.open_payroll_period(d, seed["manager1"], office_id=activity,
                               year=2026, month=10)
    pr.set_base_salary(d, seed["manager1"], office_id=activity,
                       period_id=p["period_id"], rep_user_id=seed["rep1_o1"],
                       amount=1_000_000)
    pr.grant_bonus(d, seed["supervisor1"], office_id=activity,
                   rep_user_id=seed["rep1_o1"], period_id=p["period_id"],
                   amount=50_000, reason="x")
    r = rep.monthly_report(d, seed["manager1"], office_id=activity,
                           year=2026, month=10)
    by_rep = {x["rep_id"]: x for x in r["rows"]}
    assert by_rep[seed["rep1_o1"]]["visits"] == 1
    assert by_rep[seed["rep1_o1"]]["net_salary"] == 1_050_000


def test_rep_performance(seed, activity):
    d = seed["data_dir"]
    r = rep.rep_performance(d, seed["supervisor1"], office_id=activity,
                            rep_user_id=seed["rep1_o1"],
                            start_date="2026-10-12", end_date="2026-10-18")
    row = r["rows"][0]
    assert row["visits"] == 1 and row["visits_done"] == 1
    assert row["days_present"] == 1
    # rep cannot view another rep's performance
    with pytest.raises(PermissionDenied):
        rep.rep_performance(d, seed["rep1_o1"], office_id=activity,
                            rep_user_id=seed["rep2_o1"],
                            start_date="2026-10-12", end_date="2026-10-18")


def test_reports_isolated_per_office(seed, activity):
    d = seed["data_dir"]
    r = rep.daily_report(d, seed["supervisor2"], office_id=seed["office2"],
                         day=DAY)
    # office 2 has its own rep only; office 1's activity never leaks
    assert [x["rep_id"] for x in r["rows"]] == [seed["rep_o2"]]
    assert r["totals"]["visits"] == 0
    with pytest.raises(PermissionDenied):
        rep.daily_report(d, seed["rep1_o1"], office_id=seed["office2"], day=DAY)


def test_export_csv_and_excel(seed, activity, tmp_path):
    d = seed["data_dir"]
    r = rep.daily_report(d, seed["supervisor1"], office_id=activity, day=DAY)
    csv_path = rep.export_csv(r, tmp_path / "daily.csv")
    excel_path = rep.export_excel(r, tmp_path / "daily.xlsx")
    assert csv_path.exists() and excel_path.exists()
    # round-trip: excel contains the data rows
    wb = load_workbook(excel_path)
    ws = wb.active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    flat = " ".join(str(c) for row in values for c in row if c)
    assert "التقرير اليومي" in flat
    assert str(seed["rep1_o1"]) in flat
    # csv has header + 2 data rows
    lines = (tmp_path / "daily.csv").read_text(encoding="utf-8-sig").splitlines()
    assert len(lines) == 3 and "rep_id" in lines[0]


def test_export_empty_report(tmp_path):
    p = rep.export_csv({"rows": []}, tmp_path / "empty.csv")
    assert p.exists()
